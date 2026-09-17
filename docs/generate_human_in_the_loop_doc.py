"""Generate docs/human-in-the-loop-responsibilities.docx.

Run from the repo root:

    pip install -r requirements.txt
    python docs/generate_human_in_the_loop_doc.py
"""

from pathlib import Path

from docx import Document

OUTPUT = Path(__file__).with_name("human-in-the-loop-responsibilities.docx")

TITLE = "Human-in-the-Loop Responsibilities: VM Upgrade Program with Devin"

INTRO = (
    "Devin automates the core engineering work of the VM upgrade program: "
    "analysis of each server and application, config-as-code changes, "
    "execution of upgrade scripts and playbooks, test authoring, and pull "
    "request generation. The tasks listed below remain with humans. This "
    "boundary follows the same human/automation principle already encoded in "
    "the migration playbooks: the \"Out of scope\" section of "
    "migration-playbook.md and the sign-offs in pre-migration-checklist.md."
)

SECTIONS: list[tuple[str, list[str]]] = [
    (
        "Approvals, ownership & scope decisions",
        [
            "Confirming each server/app is still in use and in scope, and who "
            "owns it (business-owner confirmation required before work starts).",
            "Change Advisory Board (CAB) / change-management approvals and "
            "maintenance-window scheduling.",
            "Defining target architecture and acceptance criteria "
            "(functional-equivalence definition).",
        ],
    ),
    (
        "Privileged infrastructure / physical / portal actions",
        [
            "Provisioning new VMs, snapshots/backups, and rollback execution at "
            "the hypervisor/cloud level.",
            "Firewall, load-balancer, DNS, and network changes.",
            "License procurement and vendor entitlement (Oracle, WAS, Pega, "
            "Confluent).",
            "Secrets/credential rotation and privileged production access.",
        ],
    ),
    (
        "Config/customization that isn't code",
        [
            "Appliance/GUI-driven middleware config: Pega rules, WAS/JBoss admin "
            "console tuning, Kafka broker/cluster config, MQ channel config. This "
            "is analogous to SAP IMG configuration, pricing procedures, output "
            "determination, workflow rules, and authorization/role design being "
            "left to functional consultants.",
        ],
    ),
    (
        "Review, validation & release gates",
        [
            "Reviewing and merging PRs (required human gate, even for Devin "
            "Review output).",
            "Release management, environment promotion, and go/no-go decisions.",
            "UAT / business acceptance testing and sign-off with app teams.",
            "Production smoke testing and post-cutover validation requiring real "
            "prod data/traffic.",
        ],
    ),
    (
        "Coordination & human judgment",
        [
            "Coordinating with app teams, downstream/integration owners, and "
            "vendors.",
            "Prioritizing/sequencing the fleet and managing interdependencies.",
            "Incident handling and rollback decisions on failed prod upgrades.",
            "Refining the playbooks themselves based on early results.",
        ],
    ),
    (
        "Things needing test data / real environments Devin can't reach",
        [
            "Providing representative test data and expected outputs.",
            "Performance/load validation against production-scale volumes.",
            "Hardware-, firmware-, or agent-level upgrades requiring "
            "console/out-of-band access.",
        ],
    ),
    (
        "Implications for the cost estimate",
        [
            "Devin ACUs cover the automatable core, not the human effort above. "
            "Total program cost = Devin ACU cost PLUS a separate human-effort "
            "line item.",
            "Per-server human overhead (review, CAB, UAT, coordination) is "
            "largely fixed regardless of automation level, so model it as a "
            "separate multiplier on top of the ACU estimate rather than folding "
            "it in.",
        ],
    ),
]


def build_document() -> Document:
    doc = Document()
    doc.add_heading(TITLE, level=1)
    doc.add_paragraph(INTRO)
    for heading, bullets in SECTIONS:
        doc.add_heading(heading, level=2)
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")
    return doc


def main() -> None:
    build_document().save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
